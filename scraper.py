import requests
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

print("[INFO] Fetching product data...")
url = "https://webscraper.io/test-sites/e-commerce/allinone/computers/laptops"
response = requests.get(url)
soup = BeautifulSoup(response.text, 'html.parser')

products = []
for item in soup.select('.thumbnail'):
    title = item.select_one('.title').get('title', '').strip() if item.select_one('.title') else item.select_one('.title').text.strip()
    price = item.select_one('.price').text.strip() if item.select_one('.price') else ''
    description = item.select_one('.description').text.strip() if item.select_one('.description') else ''
    reviews = item.select_one('.ratings .pull-right').text.strip() if item.select_one('.ratings .pull-right') else '0 reviews'
    rating = len(item.select('.ratings .glyphicon-star'))

    products.append({
        "Product Name": title,
        "Price": price,
        "Description": description,
        "Rating (Stars)": f"★ {rating}",
        "Reviews": reviews
    })

df = pd.DataFrame(products)
excel_file = "ecommerce_products_styled.xlsx"

with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
    df.to_excel(writer, index=False, sheet_name='Products')
    workbook = writer.book
    worksheet = writer.sheets['Products']

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_num in range(2, len(df) + 2):
        for col_num in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=row_num, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            if col_num in [2, 4, 5]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    col_widths = {'A': 30, 'B': 12, 'C': 50, 'D': 15, 'E': 15}
    for col_letter, width in col_widths.items():
        worksheet.column_dimensions[col_letter].width = width

    worksheet.row_dimensions[1].height = 25
    for r in range(2, len(df) + 2):
        worksheet.row_dimensions[r].height = 20

print(f"[SUCCESS] Scraped {len(products)} products! Saved to '{excel_file}'.")
